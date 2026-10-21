"""Cinq dossiers fictifs A–E, rejouables, sans créer ni changer de compte.

python -m app.seed_operations_demo --admin-email <compte-interne-existant>
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import secrets
from email.message import EmailMessage
from sqlalchemy import select
from openpyxl import Workbook
from .core import SessionLocal
from .models import Document, DocumentVersion, Folder, Launch, User
from .completeness_models import CompletenessProfile, LaunchCompleteness, RequirementRule
from .internal_models import ExperienceNote, InternalComment, LaunchOperations
from .workflow_models import DocumentReview, ImpactRule
from .quality_models import CrossDocumentRule, DocumentDatum
from .mail_models import ArchivedEmail, EmailCommitment

PROFILE_NAME = "DÉMONSTRATION A–E · exigences fictives"
SCENARIOS = {
    "A": ("FR-DEMO-2026-0141", "Chemise Oxford", "expedition", None),
    "B": ("FR-DEMO-2026-0142", "Veste coton", "fabrication", "Calcul du besoin"),
    "C": ("FR-DEMO-2026-0143", "Pantalon lin", "expedition", "Pièce douanière"),
    "D": ("FR-DEMO-2026-0144", "Surchemise", "fabrication", "pending_price"),
    "E": ("FR-DEMO-2026-0145", "Parka atelier", "fabrication", "new_version"),
}
RULES = [("Dossier technique client", "jaune", "preparation", False),
         ("Fiche d’emploi", "jaune", "preparation", True),
         ("Calcul du besoin", "jaune", "preparation", False),
         ("Fiche de prix", "jaune", "preparation", True),
         ("Bon de réception", "rose", "approvisionnement", False),
         ("Pièce douanière", "bleu", "expedition", False),
         ("Facture", "vert", "cloture", False)]


def add_version(db, user, doc, blob, filename):
    from .main import STORAGE, MIMES, audit
    from .quality_routes import auto_inspect_excel
    from pathlib import Path
    disk = secrets.token_hex(24)
    (STORAGE / disk).write_bytes(blob)
    number = max((v.number for v in doc.versions), default=0) + 1
    version = DocumentVersion(document_id=doc.id, number=number, original_name=filename,
                              disk_name=disk, size=len(blob), sha256=hashlib.sha256(blob).hexdigest(),
                              content_type=MIMES[Path(filename).suffix], uploaded_by=user.id)
    doc.versions.append(version)
    db.flush()
    auto_inspect_excel(db, version, user.id)
    audit(db, user, "document.upload" if number == 1 else "document.version.add", "document", doc.id,
          {"version": number, "sha256": version.sha256, "demo": True})
    return version


def add_document(db, user, launch, folder, kind, text):
    doc = Document(folder_id=next(f.id for f in launch.folders if f.slug == folder),
                   title=kind, doc_type=kind, created_by=user.id)
    db.add(doc)
    db.flush()
    version = add_version(db, user, doc, text.encode(), kind + ".txt")
    return doc, version


def workbook(value):
    book = Workbook()
    sheet = book.active
    sheet.title = "Besoins"
    sheet["A1"] = 1800
    sheet["A2"] = 1.8
    sheet["B1"] = value
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def seed(db, user):
    from .main import FOLDERS, audit, is_admin
    from .workflow import emit_document_impacts
    if not is_admin(user):
        raise ValueError("Un compte Administrateur existant est requis")
    profile = db.scalar(select(CompletenessProfile).where(CompletenessProfile.name == PROFILE_NAME))
    if not profile:
        profile = CompletenessProfile(name=PROFILE_NAME, description="Cas fictifs pour la recette ; règles à valider par FRAMATEX.", created_by=user.id)
        db.add(profile)
        db.flush()
        for title, folder, phase, validation in RULES:
            db.add(RequirementRule(profile_id=profile.id, title=title, folder_slug=folder,
                    document_type=title, due_phase=phase, requires_validation=validation,
                    guidance="Démonstration : contrôle humain requis."))
    result = {}
    for code, (reference, article, phase, special) in SCENARIOS.items():
        launch = db.scalar(select(Launch).where(Launch.order_ref == reference))
        if launch:
            result[code] = launch.id
            continue
        launch = Launch(name=f"Démo {code} · {article}", order_ref=reference,
                        client_name="Collection Atlas — client fictif", status="En cours", created_by=user.id)
        db.add(launch)
        db.flush()
        for slug, label in FOLDERS:
            db.add(Folder(launch_id=launch.id, slug=slug, label=label))
        db.flush()
        db.add(LaunchCompleteness(launch_id=launch.id, profile_id=profile.id, phase=phase))
        db.add(LaunchOperations(launch_id=launch.id, article=article, quantity=1800,
                provisional_ref=f"LP-DEMO-{code}", definitive_ref=f"LD-DEMO-{code}", responsible_id=user.id))
        documents = {}
        for kind, folder, _, validation in RULES:
            if kind == special:
                continue
            text = f"DONNÉES FICTIVES · {reference}\n{kind}\nQuantité : 1800 pièces\nDate : 21/10/2026\nRéférence tissu : COT-042\n"
            if kind == "Calcul du besoin":
                text += "Calcul du besoin : 3240 mètres\n"
            doc, version = add_document(db, user, launch, folder, kind, text)
            documents[kind] = doc
            if validation and not (special == "pending_price" and kind == "Fiche de prix"):
                db.add(DocumentReview(document_version_id=version.id, decision="approved",
                    comment="Validation fictive du scénario de démonstration", reviewed_by=user.id))
        if special == "new_version":
            rule = db.scalar(select(ImpactRule).where(ImpactRule.name == "DÉMO A–E · Fiche d’emploi → besoins et prix"))
            if not rule:
                rule = ImpactRule(name="DÉMO A–E · Fiche d’emploi → besoins et prix", source_kind="document",
                        source_key="Fiche d’emploi", source_folder="jaune", target_folder="jaune",
                        instruction="Examiner le calcul du besoin et la fiche de prix après la révision de la fiche d’emploi.", created_by=user.id)
                db.add(rule)
                db.flush()
            doc = documents["Fiche d’emploi"]
            version = add_version(db, user, doc, b"Revision fictive : consommation 1.9 m par piece", "fiche-emploi-v2.txt")
            emit_document_impacts(db, doc, version)
            db.add(InternalComment(document_version_id=version.id, author_id=user.id,
                                   body="Vérifier les besoins avant prise en compte de cette révision fictive."))
            for kind, quantity in (("Lancement provisoire", "1800"), ("Lancement définitif", "1750")):
                _, v = add_document(db, user, launch, "jaune", kind, f"Exemple fictif {kind} : {quantity} pièces")
                db.add(DocumentDatum(document_version_id=v.id, field_key="quantite_lancement",
                                      value_text=quantity, unit="pièce", recorded_by=user.id))
            if not db.scalar(select(CrossDocumentRule).where(CrossDocumentRule.name == "DÉMO A–E · Quantités provisoire / définitive")):
                db.add(CrossDocumentRule(name="DÉMO A–E · Quantités provisoire / définitive",
                        field_key="quantite_lancement", left_doc_type="Lancement provisoire",
                        right_doc_type="Lancement définitif", left_folder="jaune", right_folder="jaune",
                        tolerance_text="0", created_by=user.id))
            excel = Document(folder_id=next(f.id for f in launch.folders if f.slug == "jaune"),
                             title="Excel BE · calcul fictif", doc_type="Excel BE", created_by=user.id)
            db.add(excel)
            db.flush()
            add_version(db, user, excel, workbook("=A1*A2"), "besoin-v1.xlsx")
            add_version(db, user, excel, workbook(3240), "besoin-v2.xlsx")
            email = EmailMessage()
            email["From"] = "contact@client-fictif.example"
            email["To"] = "production@framatex-fictif.example"
            email["Subject"] = "Confirmation de livraison — démonstration"
            email["Date"] = "Wed, 21 Oct 2026 09:00:00 +0100"
            email.set_content("Nous vous transmettrons la confirmation le 28/10/2026, sous réserve du contrôle qualité.")
            doc = Document(folder_id=next(f.id for f in launch.folders if f.slug == "rose"),
                           title="Email · confirmation fictive", doc_type="Courriel", created_by=user.id)
            db.add(doc)
            db.flush()
            v = add_version(db, user, doc, email.as_bytes(), "confirmation.eml")
            from .mail_analysis import parse_email
            parsed = parse_email(email.as_bytes())
            archived = ArchivedEmail(launch_id=launch.id, document_version_id=v.id, archived_by=user.id,
                attachments_json=json.dumps(parsed["attachments"]), **{k: parsed[k] for k in
                    ("message_id", "subject", "sender", "recipients", "carbon_copy", "sent_at_raw", "body_excerpt")})
            db.add(archived)
            db.flush()
            db.add(EmailCommitment(launch_id=launch.id, email_id=archived.id,
                    description="Confirmer la livraison après contrôle qualité (démonstration)",
                    responsible="Contact client fictif", due_date="2026-10-28", condition="Sous réserve du contrôle qualité",
                    source_excerpt=parsed["body_excerpt"], confirmed_by=user.id))
            db.add(ExperienceNote(launch_id=launch.id, party_type="fournisseur", party_name="Tissage Atlas (fictif)",
                                  category="documentaire", body="Exiger la référence du lot sur le bon de livraison.", author_id=user.id))
        audit(db, user, "demo.launch.create", "launch", launch.id, {"scenario": code, "demo": True})
        result[code] = launch.id
    db.commit()
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--admin-email", required=True)
    args = parser.parse_args()
    from .migrate import migrate
    migrate()
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.email == args.admin_email.strip().lower()))
        if not user:
            parser.error("Compte interne existant introuvable ; aucun compte n'a été créé")
        result = seed(db, user)
        print(f"{len(result)} dossiers de démonstration A–E disponibles ; aucun compte modifié.")


if __name__ == "__main__":
    main()
