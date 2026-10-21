"""Comparaisons explicites et inspection statique Excel sans exécuter de formules.

Les constats sont des pistes de vérification, jamais des non-conformités automatiques.
"""
from __future__ import annotations
import hashlib
import json
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from .models import Document, DocumentVersion, Folder
from .quality_models import DocumentDatum, CrossDocumentRule, QualityDecision

MAX_XLSX_BYTES = 5 * 1024 * 1024
MAX_CELLS = 40000
MAX_FINDINGS = 200


def latest_document_versions(db: Session, launch_id: str, doc_type: str, folder: str):
    docs = db.scalars(select(Document).join(Folder).where(
        Folder.launch_id == launch_id, Folder.slug == folder, Document.doc_type == doc_type
    )).all()
    return [(d, max(d.versions, key=lambda x:x.number)) for d in docs if d.versions]


def comparison(db: Session, launch_id: str, rule: CrossDocumentRule):
    """Do not guess when several documents could be the expected source."""
    left = latest_document_versions(db, launch_id, rule.left_doc_type, rule.left_folder)
    right = latest_document_versions(db, launch_id, rule.right_doc_type, rule.right_folder)
    status = 'unavailable'; reason = 'Pièce manquante'; pair = []
    if not left or not right:
        reason = 'Pièce requise absente' if not left or not right else reason
    elif len(left) != 1 or len(right) != 1 or left[0][0].id == right[0][0].id:
        status = 'ambiguous'; reason = 'Plusieurs documents candidats : sélectionner une source métier précise'
    else:
        for d, v in [left[0], right[0]]:
            datum = db.scalar(select(DocumentDatum).where(
                DocumentDatum.document_version_id == v.id, DocumentDatum.field_key == rule.field_key
            ))
            pair.append({'document_id': d.id, 'document_title': d.title, 'version_id': v.id,
                         'version': v.number, 'folder': d.folder.slug, 'value': datum.value_text if datum else None,
                         'unit': datum.unit if datum else None})
        if not all(x['value'] is not None for x in pair):
            status = 'unavailable'; reason = 'Valeur non relevée sur une ou plusieurs versions courantes'
        elif pair[0]['unit'] != pair[1]['unit']:
            status = 'unit_mismatch'; reason = 'Unités incompatibles : conversion manuelle requise'
        else:
            delta = abs(Decimal(pair[0]['value'])-Decimal(pair[1]['value']))
            status = 'consistent' if delta <= Decimal(rule.tolerance_text) else 'difference'
            reason = 'Écart dans la tolérance' if status == 'consistent' else 'Écart à examiner'
    source = {'rule':rule.id,'pairs':[(p['version_id'],p['value'],p['unit']) for p in pair]}
    signature = hashlib.sha256(json.dumps(source,sort_keys=True).encode()).hexdigest()
    return {'rule_id':rule.id,'name':rule.name, 'field_key':rule.field_key,
            'left_doc_type':rule.left_doc_type, 'right_doc_type':rule.right_doc_type,
            'folders':[rule.left_folder,rule.right_folder], 'tolerance':rule.tolerance_text,
            'status':status,'reason':reason,'pair':pair,'signature':signature}


def _read_cells(path: Path):
    if path.stat().st_size > MAX_XLSX_BYTES:
        raise ValueError('Analyse limitée aux classeurs de 5 Mo maximum')
    # Les fichiers sont conservés sous un nom opaque, sans extension : fournir un flux.
    stream = path.open("rb")
    try:
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
    except Exception:
        stream.close()
        raise
    data = {}; seen = 0; formula_count=0
    try:
        for sheet in book:
            if sheet.max_row * sheet.max_column > MAX_CELLS * 10:
                raise ValueError('Classeur trop volumineux : inspection manuelle nécessaire')
            for row in sheet.iter_rows():
                for cell in row:
                    seen+=1
                    if seen > MAX_CELLS:
                        raise ValueError('Plus de 40 000 cellules : inspection non réalisée')
                    v = cell.value
                    if v is None: continue
                    is_formula = cell.data_type == 'f' or (isinstance(v,str) and v.startswith('='))
                    if is_formula: formula_count+=1
                    data[(sheet.title,cell.coordinate)] = (is_formula, str(v)[:400] if is_formula else str(v)[:120])
    finally:
        book.close()
        stream.close()
    return data,seen,formula_count


def analyze_workbook(path: Path, baseline: Path | None):
    """Static formula-integrity checks only; does not calculate or execute Excel."""
    try:
        current, scanned, formulas = _read_cells(path)
        before,_,_ = _read_cells(baseline) if baseline else ({},0,0)
    except Exception as ex:
        return {'status':'limited','scanned_cells':0,'formula_count':0,
                'findings':[{'type':'inspection_limit','level':'warning','sheet':'','cell':'',
                             'message':str(ex)[:200]}]}
    findings=[]
    def add(kind, level, sheet, coord, message, before=None, after=None):
        if len(findings) < MAX_FINDINGS:
            findings.append({'type':kind,'level':level,'sheet':sheet,'cell':coord,'message':message,
                             'before':before,'after':after})
    for (sheet,coord),(is_formula, value) in current.items():
        if is_formula and '#REF!' in value.upper():
            add('broken_reference','critical',sheet,coord,'Référence #REF! détectée dans la formule', after=value)
    if baseline:
        for (sheet,coord),(was_formula, prev) in before.items():
            if not was_formula:continue
            now=current.get((sheet,coord))
            if now is None or not now[0]:
                add('formula_removed','critical',sheet,coord,'Formule remplacée ou supprimée depuis la version précédente', prev, now[1] if now else None)
            elif now[1] != prev:
                add('formula_changed','warning',sheet,coord,'Formule modifiée : vérifier la logique du calcul', prev, now[1])
        status='reviewed'
    else:
        status='first_version'
    if len(findings)==MAX_FINDINGS:
        add('truncated','warning','','','Affichage plafonné à 200 constats')
    return {'status':status,'scanned_cells':scanned,'formula_count':formulas,'findings':findings}
