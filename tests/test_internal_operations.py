"""Recette API du pilotage interne et non-régression des frontières Client."""
import hashlib
import io
import uuid
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, inspect, select, text
from app.core import Base, SessionLocal, hash_password
from app.main import app
from app.models import Document, Launch, Role, RolePermission, User
from app.internal_models import InternalComment
from app.seed_operations_demo import seed

H = {"X-Framatex-Request": "1"}
PASSWORD = "InternalRecipePassword2026!"


@pytest.fixture(scope="module")
def setup():
    with TestClient(app) as api:
        suffix = uuid.uuid4().hex[:8]
        with SessionLocal() as db:
            role = db.scalar(select(Role).where(Role.name == "Administrateur"))
            if not role:
                role = Role(name="Administrateur")
                db.add(role)
            user = User(email=f"ops-{suffix}@example.test", name="Responsable recette",
                        password_hash=hash_password(PASSWORD), roles=[role])
            worker_role = Role(name=f"Réception recette {suffix}", grants=[
                RolePermission(action="launch.read"), RolePermission(action="audit.read"),
                RolePermission(action="document.read", folder_slug="rose"),
                RolePermission(action="document.write", folder_slug="rose")])
            worker = User(email=f"rose-{suffix}@example.test", name="Réception recette",
                          password_hash=hash_password(PASSWORD), roles=[worker_role])
            db.add_all([user, worker])
            db.commit()
            launches = seed(db, user)
            second_run = seed(db, user)
            assert launches == second_run
            user_email, worker_email, user_id = user.email, worker.email, user.id
        assert api.post('/api/auth/login', headers=H,
                        json={"email": user_email, "password": PASSWORD}).status_code == 200
        yield {"api": api, "launches": launches, "email": user_email, "worker_email": worker_email,
               "user_id": user_id, "suffix": suffix}


def scratch(api, name="Recette indépendante"):
    response = api.post('/api/launches', headers=H, json={"name": name})
    assert response.status_code == 201
    return response.json()['id']


def upload(api, lid, blob=b"Document recette", name="recette.txt", folder="jaune", title="Document recette", **data):
    return api.post(f'/api/launches/{lid}/documents', headers=H,
                    data={"folder": folder, "title": title, **data}, files={"file": (name, blob)})


def test_dashboard_missing_due_states_and_demo_are_real(setup):
    api, ids = setup['api'], setup['launches']
    report = api.get('/api/workspace/dashboard')
    assert report.status_code == 200, report.text
    body = report.json()
    assert set(body['kpis']) >= {'active_launches', 'complete_launches', 'incomplete_launches',
        'missing_documents', 'pending_documents', 'alerts', 'inconsistencies', 'unacknowledged_changes'}
    rows = {r['id']: r for r in body['launches']}
    assert rows[ids['A']]['completion']['state'] == 'complete'
    assert rows[ids['A']]['completion']['percent'] == 100
    assert rows[ids['B']]['completion']['state'] == 'incomplete'
    missing = {(r['launch_id'], r['title']) for r in body['missing_documents']}
    assert (ids['B'], 'Calcul du besoin') in missing
    assert (ids['C'], 'Pièce douanière') in missing
    assert all(r['status'] == 'missing' and r['due'] for r in body['missing_documents'])
    assert not any(title == 'Facture' for _, title in missing)
    d = api.get(f"/api/launches/{ids['D']}/completeness").json()
    assert any(r['document_type'] == 'Fiche de prix' and r['status'] == 'pending_validation' for r in d['rows'])
    assert rows[ids['E']]['unacknowledged_changes'] >= 1
    assert len(rows[ids['E']]['inconsistencies']) >= 1
    assert sum(f.get('pending_validation', 0) for f in rows[ids['A']]['folders']) == 0
    assert sum(f.get('pending_validation', 0) for f in rows[ids['D']]['folders']) == 1


def test_actions_target_folder_document_and_quality(setup):
    body = setup['api'].get('/api/workspace/priorities').json()
    missing = next(r for r in body['items'] if r['launch_id'] == setup['launches']['B'] and r['kind'] == 'missing')
    assert missing['folder'] == 'jaune'
    validation = next(r for r in body['items'] if r['launch_id'] == setup['launches']['D'] and r['kind'] == 'validation')
    assert validation['document_id']
    assert any(r['kind'] == 'quality' for r in body['items'])
    assert any(r['kind'] == 'excel' for r in body['items'])


def test_launch_tracking_and_filters(setup):
    api, lid = setup['api'], setup['launches']['B']
    result = api.patch(f'/api/launches/{lid}/operations', headers=H, json={
        'article': 'Veste coton recette', 'quantity': 2400, 'provisional_ref': 'LP-RECETTE',
        'definitive_ref': 'LD-RECETTE', 'responsible_id': setup['user_id']})
    assert result.status_code == 200, result.text
    assert result.json()['quantity'] == 2400
    rows = api.get('/api/workspace/launches', params={'q': 'LP-RECETTE', 'article': 'coton',
             'completion': 'incomplete', 'status': 'En cours', 'date_from': '2020-01-01'}).json()
    assert [r['id'] for r in rows] == [lid]
    assert api.get('/api/workspace/launches', params={'q': 'does-not-exist'}).json() == []
    assert api.patch(f'/api/launches/{lid}/operations', headers=H, json={'quantity': -1}).status_code == 422
    assert api.patch(f'/api/launches/{lid}/operations', headers=H, json={'responsible_id': 'missing'}).status_code == 422


def test_comments_sha_duplicates_versions_passport_and_audit(setup):
    api = setup['api']
    lid = scratch(api)
    first = upload(api, lid, comment='Premier commentaire interne', doc_type='Calcul du besoin')
    assert first.status_code == 201, first.text
    doc = first.json()
    assert doc['latest']['sha256'] == hashlib.sha256(b'Document recette').hexdigest()
    comments = api.get(f"/api/document-versions/{doc['latest']['id']}/comments").json()
    assert comments[0]['body'] == 'Premier commentaire interne'
    assert comments[0]['author'] == 'Responsable recette'
    duplicate = upload(api, lid, title='Copie métier').json()
    assert duplicate['duplicates'][0]['document_id'] == doc['id']
    same = api.post(f"/api/documents/{doc['id']}/versions", headers=H,
                    files={'file': ('recette.txt', b'Document recette')})
    assert same.status_code == 409
    revised = api.post(f"/api/documents/{doc['id']}/versions", headers=H,
                       data={'comment': 'Version corrigée et commentée'}, files={'file': ('v2.txt', b'Nouvelle version')})
    assert revised.status_code == 201
    assert len(revised.json()['versions']) == 2
    v2 = revised.json()['latest']['id']
    assert api.get(f"/api/documents/{doc['id']}/versions/1/download").content == b'Document recette'
    assert api.get(f'/api/document-versions/{v2}/passport').json()['integrity']['status'] == 'intact'
    assert api.post(f'/api/document-versions/{v2}/comments', headers=H,
                    json={'body': '<script>texte interne</script>'}).status_code == 201
    assert api.post(f'/api/document-versions/{v2}/comments', headers=H, json={'body': '  '}).status_code == 422
    events = api.get('/api/workspace/activity', params={'document_id': doc['id'], 'action': 'document.comment'}).json()
    assert events['total'] == 3
    assert all(row['document_id'] == doc['id'] and row['actor'] == 'Responsable recette' for row in events['items'])
    summary = api.get(f'/api/launches/{lid}/operations').json()
    assert summary['last_activity'] >= events['items'][0]['occurred_at']
    assert api.get('/api/workspace/activity', params={'launch_id': lid, 'date_to': '2000-01-01'}).json()['total'] == 0


@pytest.mark.parametrize('name,blob,expected', [
    ('unsafe.exe', b'not allowed', 422), ('unsafe.js', b'alert(1)', 422),
    ('false.pdf', b'not a PDF', 422), ('false.png', b'not a PNG', 422),
    ('false.docx', b'not a ZIP', 422), ('false.xlsx', b'not a ZIP', 422),
    ('empty.txt', b'', 422), ('program.txt', b'MZ executable content', 422),
    ('valid.csv', b'reference,quantite\nCOT-042,1800\n', 201),
    ('invalid.csv', b'\x00\x01,\x00', 422), ('binary.csv', b'\xff\xfe\x00', 422),
    ('old.xls', b'unsupported legacy format', 422)])
def test_upload_types_are_validated(setup, name, blob, expected):
    lid = scratch(setup['api'], 'Recette format ' + name)
    assert upload(setup['api'], lid, blob, name).status_code == expected


def test_supported_office_images_and_size_limit(setup):
    from PIL import Image
    from docx import Document as Word
    from openpyxl import Workbook
    api, lid = setup['api'], scratch(setup['api'])
    word = Word(); word.add_paragraph('Dossier technique recette')
    out = io.BytesIO(); word.save(out)
    assert upload(api, lid, out.getvalue(), 'fiche.docx', title='Word').status_code == 201
    book = Workbook(); book.active['A1'] = '=2+2'
    out = io.BytesIO(); book.save(out)
    assert upload(api, lid, out.getvalue(), 'fiche.xlsx', title='Excel').status_code == 201
    for kind in ('PNG', 'JPEG'):
        out = io.BytesIO(); Image.new('RGB', (32,32), 'white').save(out, kind)
        assert upload(api, lid, out.getvalue(), 'image.' + ('jpg' if kind == 'JPEG' else 'png'), title=kind).status_code == 201
    assert upload(api, lid, b'x' * (20 * 1024 * 1024 + 1), 'large.txt').status_code == 413


def test_search_versions_metadata_and_current_review(setup):
    api, lid = setup['api'], setup['launches']['E']
    report = api.get('/api/workspace/search', params={'q': 'emploi', 'launch_id': lid, 'folder': 'jaune'}).json()
    rows = report['items']
    assert {r['version'] for r in rows} == {1, 2}
    assert next(r for r in rows if r['version'] == 1)['status'] == 'approved'
    assert next(r for r in rows if r['version'] == 2)['status'] == 'to_review'
    assert api.get('/api/workspace/search', params={'q': 'fiche-emploi-v2.txt'}).json()['total'] >= 1
    assert api.get('/api/workspace/search', params={'date_to': '2000-01-01'}).json()['total'] == 0
    one = api.get('/api/workspace/search', params={'launch_id': lid, 'limit': 1}).json()
    assert one['total'] > 1 and len(one['items']) == 1


def test_intake_classification_keeps_calcul_du_besoin_in_jaune(setup):
    response = setup['api'].post(f"/api/launches/{setup['launches']['B']}/intake/preview", headers=H,
               files={'file': ('calcul_du_besoin.txt', 'Calcul du besoin\nQuantité : 1500 pièces'.encode())})
    assert response.status_code == 200
    proposals = response.json()['classification']['suggestions']
    assert any(p['folder'] == 'jaune' and p['doc_type'] == 'Calcul du besoin' for p in proposals)
    assert response.json()['fields']


def test_excel_before_after_and_eml_remain_available(setup):
    api, lid = setup['api'], setup['launches']['E']
    excel = api.get(f'/api/launches/{lid}/excel-inspections').json()
    findings = [f for row in excel if row['version'] == 2 for f in row['inspection']['findings']]
    assert any(f['type'] == 'formula_removed' and f['before'] == '=A1*A2' and f['after'] == '3240' for f in findings)
    emails = api.get(f'/api/launches/{lid}/emails').json()
    assert emails['emails'] and emails['commitments']
    assert emails['emails'][0]['sender'] == 'contact@client-fictif.example'


def test_internal_experience_history_and_audit(setup):
    api, lid = setup['api'], setup['launches']['B']
    for party in ('client', 'fournisseur', 'atelier'):
        response = api.post(f'/api/launches/{lid}/experience-notes', headers=H,
               json={'party_type': party, 'party_name': 'Partenaire recette ' + party,
                     'category': 'incident', 'body': 'Retour interne confidentiel ' + party})
        assert response.status_code == 201, response.text
    assert len(api.get(f'/api/launches/{lid}/experience-notes').json()) == 3
    rows = api.get('/api/workspace/experience', params={'party_type': 'atelier', 'q': 'recette'}).json()
    assert len(rows) == 1 and rows[0]['party_type'] == 'atelier'
    assert api.get('/api/workspace/activity', params={'launch_id': lid, 'action': 'experience.note.create'}).json()['total'] == 3


def test_scope_prevents_aggregate_search_comment_and_audit_leaks(setup):
    api = setup['api']
    hidden = scratch(api, 'Secret recette inaccessible')
    doc = upload(api, hidden, b'valeur secretest', folder='rose', title='SECRET-INTERNAL-RECIPE').json()
    assert api.put(f'/api/launches/{hidden}/access', headers=H, json={'restricted': True}).status_code == 200
    with TestClient(app) as worker:
        assert worker.post('/api/auth/login', headers=H, json={
            'email': setup['worker_email'], 'password': PASSWORD}).status_code == 200
        assert hidden not in {r['id'] for r in worker.get('/api/workspace/launches').json()}
        dashboard = worker.get('/api/workspace/dashboard')
        assert dashboard.status_code == 200, dashboard.text
        assert 'SECRET-INTERNAL-RECIPE' not in dashboard.text
        assert 'Calcul du besoin' not in dashboard.text
        assert worker.get('/api/workspace/search', params={'q': 'SECRET-INTERNAL-RECIPE'}).json()['total'] == 0
        assert worker.get(f"/api/document-versions/{doc['latest']['id']}/comments").status_code == 404
        assert worker.get('/api/workspace/activity', params={'launch_id': hidden}).json()['total'] == 0
        assert worker.get('/api/workspace/experience').status_code == 403
        assert worker.get('/api/admin/client-order-requests').status_code == 403
        partial = worker.get(f"/api/launches/{setup['launches']['B']}/operations").json()
        assert partial['completion']['state'] == 'partial'
        assert all('documents' not in f for f in partial['folders'] if not f['can_read'])


def test_client_inbox_downloads_and_no_internal_publication(setup):
    api = setup['api']
    with TestClient(app) as client:
        email = f"ops-client-{setup['suffix']}@example.test"
        account = client.post('/api/client/register', headers=H, json={
            'raison_sociale': 'Client recette', 'nom_contact': 'Contact Client', 'email': email,
            'telephone': '0600000000', 'adresse': 'Adresse fictive', 'pays': 'Maroc', 'password': PASSWORD}).json()
        assert api.patch(f"/api/admin/clients/{account['id']}/status", headers=H, json={'status': 'active'}).status_code == 200
        assert client.post('/api/client/login', headers=H, json={'email': email, 'password': PASSWORD}).status_code == 200
        request = client.post('/api/client/order-requests', headers=H, json={
            'article': 'Modèle client', 'quantite': 400, 'date_souhaitee': '2027-01-01', 'status': 'submitted'}).json()
        doc = client.post(f"/api/client/order-requests/{request['id']}/documents", headers=H,
              data={'title': 'Spécification client'}, files={'file': ('spec.txt', b'Specification du client')}).json()
        detail = api.get(f"/api/admin/client-order-requests/{request['id']}")
        assert detail.status_code == 200
        assert detail.json()['documents'][0]['id'] == doc['id']
        path = f"/api/admin/client-order-requests/{request['id']}/documents/{doc['id']}/versions/1/download"
        assert api.get(path).content == b'Specification du client'
        assert client.get(path).status_code == 401
        for path in ('/api/workspace/dashboard', '/api/workspace/search', '/api/workspace/activity',
                     '/api/workspace/experience', '/api/admin/clients', '/api/admin/client-order-requests'):
            assert client.get(path).status_code == 401, path
        own = client.get(f"/api/client/order-requests/{request['id']}")
        assert 'Retour interne' not in own.text and 'fiche de prix' not in own.text.lower()
        assert len(own.json()['documents']) == 1


@pytest.mark.parametrize('path', ['/', '/espace-framatex', '/espace-client', '/api/health'])
def test_principal_routes_and_security_headers(setup, path):
    response = setup['api'].get(path)
    assert response.status_code == 200
    assert response.headers['x-content-type-options'] == 'nosniff'
    if path != '/api/health':
        assert "script-src 'self'" in response.headers['content-security-policy']


def test_migration_from_previous_schema_preserves_accounts_and_is_repeatable(tmp_path, monkeypatch):
    import importlib
    migration = importlib.import_module('app.migrate')
    old_engine = create_engine('sqlite:///' + str(tmp_path / 'legacy.db'))
    new_tables = {'launch_operations', 'internal_document_comments', 'internal_experience_notes'}
    Base.metadata.create_all(old_engine, tables=[table for name, table in Base.metadata.tables.items() if name not in new_tables])
    with old_engine.begin() as conn:
        for revision in migration.REVISIONS[:-1]:
            conn.execute(text('INSERT INTO schema_migrations VALUES (:v,:d)'), {'v': revision, 'd': '2026-01-01'})
        conn.execute(text("INSERT INTO users (id,email,name,password_hash,active,created_at) VALUES ('keep','keep@example.test','Compte existant','hash-preserve',1,'2026-01-01')"))
    monkeypatch.setattr(migration, 'engine', old_engine)
    migration.migrate()
    migration.migrate()
    assert new_tables.issubset(inspect(old_engine).get_table_names())
    with old_engine.connect() as conn:
        assert conn.execute(text("SELECT password_hash FROM users WHERE id='keep'")).scalar() == 'hash-preserve'
        assert conn.execute(text("SELECT count(*) FROM schema_migrations WHERE version='0005_internal_operations'")).scalar() == 1
    assert list(tmp_path.glob('legacy.db.backup-*'))
    old_engine.dispose()
