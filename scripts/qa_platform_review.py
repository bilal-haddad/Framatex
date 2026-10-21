"""Recette fonctionnelle complète de la plateforme FRAMATEX (API réelle, base vierge).

python scripts/qa_platform_review.py --output data/qa-platform

Démarre un serveur éphémère sur une base vierge (toutes migrations), crée un
administrateur et les dossiers de démonstration A–E, puis vérifie une à une les
fonctions de la plateforme : routes, sécurité, documents et formats, SHA-256,
doublons, versions, complétude (5 états), documents manquants, alertes d'impact,
cohérence, OCR/classification (image et PDF scanné), Excel, EML, recherche,
mémoire interne, workflow Client ↔ FRAMATEX (demandes, pièces demandées,
partages), isolation A/B, droits d'un collaborateur, accès restreint, circuits de
validation, exigences, matières, simulation d'impact, import ORLI (fichier),
demandes fournisseurs, recette pilote, guide, audit, QR et migration d'une copie
de base. N'utilise aucune donnée réelle.
Résultat : <output>/qa-results.json et un résumé en console.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import secrets
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
from contextlib import contextmanager
from datetime import datetime, timezone
from types import SimpleNamespace
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
H = {'X-Framatex-Request': '1'}


class QA:
    def __init__(self):
        self.results = []
        self.secrets = set()

    def redact(self, value):
        result = str(value)
        for secret in self.secrets:
            result = result.replace(secret, '[masqué]')
        return result

    def check(self, area, name, ok, detail=''):
        detail = self.redact(detail)
        self.results.append({'area': area, 'check': name, 'ok': bool(ok), 'detail': detail[:400]})
        print(f"{'OK ' if ok else 'ÉCHEC'} [{area}] {name}{(' — ' + detail[:160]) if detail and not ok else ''}", flush=True)
        return ok

    def run(self, area, name, fn):
        try:
            value = fn()
            return self.check(area, name, value is not False, '' if value in (True, None) else value)
        except Exception as exc:  # une vérification en échec ne doit pas arrêter la recette
            return self.check(area, name, False, f'{type(exc).__name__}: {exc}')


# ---------------------------------------------------------------------------
# Fichiers de recette (générés en mémoire, contenus fictifs)
# ---------------------------------------------------------------------------
def pdf_bytes(text='Recette FRAMATEX'):
    return (b'%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n2 0 obj<</Type/Pages/Kids[]/Count 0>>endobj\n'
            b'% ' + text.encode() + b'\ntrailer<</Root 1 0 R>>\n%%EOF\n')


def docx_bytes(text):
    from docx import Document
    d = Document(); d.add_paragraph(text); out = io.BytesIO(); d.save(out); return out.getvalue()


def xlsx_bytes(cells):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = 'Calcul'
    for coord, value in cells.items():
        ws[coord] = value
    out = io.BytesIO(); wb.save(out); return out.getvalue()


def image_bytes(fmt, text=None):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new('RGB', (1400, 360) if text else (40, 40), 'white')
    if text:
        draw = ImageDraw.Draw(img)
        font = None
        for path in ('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'):
            if Path(path).exists():
                font = ImageFont.truetype(path, 54); break
        y = 40
        for line in text.split('\n'):
            draw.text((40, y), line, fill='black', font=font); y += 90
    out = io.BytesIO(); img.save(out, fmt); return out.getvalue()


def eml_bytes():
    return ('From: Atelier Recette <atelier@exemple.test>\r\nTo: FRAMATEX <contact@framatex.net>\r\n'
            'Subject: Confirmation delai lot 0143\r\nDate: Tue, 20 Oct 2026 09:15:00 +0200\r\n'
            'Message-ID: <qa-0143@exemple.test>\r\nContent-Type: text/plain; charset=utf-8\r\n\r\n'
            'Bonjour,\r\nNous confirmons la livraison du lot pour le 15/11/2026.\r\nCordialement.\r\n').encode()


def sha(data):
    return hashlib.sha256(data).hexdigest()


# ---------------------------------------------------------------------------
def start_server(env):
    with socket.socket() as sock:
        sock.bind(('127.0.0.1', 0)); port = sock.getsockname()[1]
    with (Path(env['FRAMATEX_DATA_DIR']) / 'server.log').open('w', encoding='utf-8') as log:
        server = subprocess.Popen([sys.executable, str(ROOT / 'scripts' / 'windows_runtime.py'), 'serve', '--port', str(port),
                                   '--demo-email', 'admin.test@framatex.com', '--no-browser'],
                                  cwd=ROOT, env=env, stdout=log, stderr=log)
    import httpx
    base = f'http://127.0.0.1:{port}'
    for _ in range(450):
        try:
            runtime = json.loads((Path(env['FRAMATEX_DATA_DIR']) / 'runtime.json').read_text())
            base = runtime['url'].rstrip('/')
            response = httpx.get(base + '/api/health', timeout=1, trust_env=False)
            if response.status_code == 200 and response.json().get('runtime_id') == runtime['runtime_id']:
                return server, base
        except (httpx.TransportError, OSError, ValueError):
            pass
        if server.poll() is not None:
            raise RuntimeError('Serveur de recette arrêté : ' + (Path(env['FRAMATEX_DATA_DIR']) / 'server.log').read_text()[-4000:])
        time.sleep(.1)
    server.terminate(); server.wait(timeout=15)
    raise RuntimeError('Serveur de recette non prêt')


@contextmanager
def platform_server(qa):
    """Base vierge migrée, administrateur et dossiers A–E, serveur Uvicorn réel.

    Les mots de passe sont générés en mémoire pour la durée de la recette ; aucun
    n'est écrit dans le dépôt ni dans le rapport. Après l'arrêt du serveur, la
    migration est rejouée sur une copie de la base pour vérifier qu'elle est additive.
    """
    with tempfile.TemporaryDirectory(prefix='framat-qa-') as tmp:
        env = {**os.environ, 'FRAMATEX_DATA_DIR': tmp, 'FRAMATEX_SECRET_KEY': secrets.token_urlsafe(48)}
        env.pop('DATABASE_URL', None); env.pop('FRAMATEX_COOKIE_SECURE', None)
        os.environ.update({k: env[k] for k in ('FRAMATEX_DATA_DIR', 'FRAMATEX_SECRET_KEY')})
        sys.path.insert(0, str(ROOT))
        from app.migrate import migrate
        from app.core import SessionLocal, hash_password
        from app.models import Role, User
        from app.seed_operations_demo import seed
        migrate()
        qa.check('Base', 'Base vierge : toutes les migrations appliquées', True)
        password = secrets.token_urlsafe(18) + 'Aa1!'
        qa.secrets.add(password)
        with SessionLocal() as db:
            role = Role(name='Administrateur')
            db.add(User(name='Responsable recette', email='recette@framatex.test', password_hash=hash_password(password), roles=[role]))
            # Same target email as -Demo, but only in this disposable database.
            # No real PC account/password is inspected or changed.
            db.add(User(name='Administrateur de test', email='admin.test@framatex.com', password_hash=hash_password(password), roles=[role]))
            db.commit()
            admin_user = db.query(User).filter_by(email='recette@framatex.test').one()
            ids = seed(db, admin_user)
        env['FRAMATEX_SECRET_INPUT'] = password
        server, base = start_server(env)
        env.pop('FRAMATEX_SECRET_INPUT', None)
        try:
            yield SimpleNamespace(base=base, password=password, ids=ids, tmp=tmp, env=env, server=server)
        finally:
            server.terminate()
            try:
                server.wait(timeout=10)
            except subprocess.TimeoutExpired:
                server.kill()
                server.wait(timeout=5)
        run_migration_copy(qa, tmp, env)


def write_report(qa, output, name='qa-results.json'):
    passed = sum(r['ok'] for r in qa.results)
    report = {'tested_at': datetime.now(timezone.utc).isoformat(),
              'source_revision': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
              'source_dirty': bool(subprocess.check_output(['git', 'status', '--porcelain'], cwd=ROOT, text=True).strip()),
              'passed': passed, 'failed': len(qa.results) - passed, 'checks': qa.results,
              'runtime': getattr(qa, 'runtime', None)}
    output.mkdir(parents=True, exist_ok=True)
    (output / name).write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"\n{passed}/{len(qa.results)} vérifications réussies.", flush=True)
    return 0 if passed == len(qa.results) else 1


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT / 'data' / 'qa-platform')
    args = parser.parse_args()
    qa = QA()
    with platform_server(qa) as ctx:
        run_checks(qa, ctx.base, ctx.password, ctx.ids, ctx.tmp)
        check_runtime_liveness(qa, ctx)
    sys.exit(write_report(qa, args.output.resolve()))


def run_migration_copy(qa, tmp, env):
    src = Path(tmp) / 'framatex.db'
    if not src.exists():
        candidates = list(Path(tmp).glob('*.db')) + list(Path(tmp).glob('*.sqlite*'))
        src = candidates[0] if candidates else None
    if not src:
        qa.check('Base', 'Copie de base existante : migration additive', False, 'fichier SQLite introuvable'); return
    with tempfile.TemporaryDirectory(prefix='framat-qa-copy-') as copy_dir:
        dst = Path(copy_dir) / src.name
        shutil.copy2(src, dst)
        tables = ['users', 'launches', 'documents', 'document_versions', 'audit_events', 'client_accounts',
                  'client_order_requests', 'client_document_requests', 'client_shared_documents']
        def counts(path):
            with sqlite3.connect(path) as conn:
                return {t: conn.execute(f'SELECT count(*) FROM {t}').fetchone()[0] for t in tables}
        before = counts(dst)
        copy_env = {**env, 'FRAMATEX_DATA_DIR': copy_dir}
        upgrade = subprocess.run([sys.executable, '-m', 'app.migrate', 'upgrade'], cwd=ROOT, env=copy_env, capture_output=True, text=True)
        again = subprocess.run([sys.executable, '-m', 'app.migrate', 'upgrade'], cwd=ROOT, env=copy_env, capture_output=True, text=True)
        after = counts(dst)
        qa.check('Base', 'Copie de base existante : migration rejouable sans erreur', upgrade.returncode == 0 and again.returncode == 0,
                 (upgrade.stderr or again.stderr)[-300:])
        qa.check('Base', 'Copie de base existante : aucune perte de données', before == after, {'avant': before, 'après': after})


def check_runtime_liveness(qa, ctx):
    """Require all five probes, then an independent post-browser HTTP check.
    Called while still inside platform_server: its cleanup is not readiness.
    """
    import httpx
    path = Path(ctx.tmp) / 'runtime.json'
    deadline = time.monotonic() + 75
    report = {}
    while time.monotonic() < deadline and ctx.server.poll() is None:
        try:
            report = json.loads(path.read_text(encoding='utf-8'))
        except (OSError, ValueError):
            pass
        if report.get('status') in ('ready', 'degraded'):
            break
        time.sleep(.2)
    checks = report.get('health', [])
    for offset in (0, 2, 10, 30, 60):
        sample = next((s for s in checks if s['at'] == offset), {})
        qa.check('Runtime', f'T+{offset}s : même PID, LISTENING et /api/health 200 rapide',
                 sample.get('ok') is True and sample.get('pid') == report.get('pid') and sample.get('status') == 200
                 and sample.get('ms', 99999) < 3000 and sample.get('listening') is True, sample)
    qa.check('Runtime', 'Login admin.test@framatex.com jetable et Chromium réels : aucun faux succès ni arrêt après sondes',
             report.get('status') == 'ready' and not report.get('errors') and ctx.server.poll() is None
             and any('admin.test@framatex.com' in check for check in report.get('checks', [])), report.get('errors'))
    def alive_after_browser():
        with httpx.Client(timeout=3, trust_env=False) as client:
            value = client.get(ctx.base + '/api/health').json()
            assert value.get('pid') == report['pid'] and value.get('runtime_id') == report['runtime_id']
            for route in ('/', '/espace-client', '/espace-framatex'):
                assert client.get(ctx.base + route).status_code == 200, route
        assert ctx.server.poll() is None
    qa.run('Runtime', 'Après fermeture du navigateur de test : serveur et trois pages toujours disponibles', alive_after_browser)
    # Keep evidence outside the deleted temporary DB without embedding secrets.
    qa.runtime = report


def run_checks(qa, base, password, ids, tmp):
    import httpx
    admin = httpx.Client(base_url=base, headers=H, timeout=60, trust_env=False)
    anon = httpx.Client(base_url=base, timeout=30, trust_env=False)

    # ------------------------------------------------------------- Routes & landing
    def routes():
        for path in ('/', '/espace-client', '/espace-framatex', '/api/health'):
            r = anon.get(path); assert r.status_code == 200, (path, r.status_code)
            if path != '/api/health':
                assert "frame-ancestors 'none'" in r.headers['content-security-policy'], path
        html = anon.get('/').text
        assert 'href="/espace-client"' in html and 'href="/espace-framatex"' in html
    qa.run('Landing', 'Routes /, /espace-client, /espace-framatex, /api/health et CSP', routes)
    def assets():
        for path in ('/assets/landing/logo-framatex.jpg', '/assets/brand/tokens.css', '/assets/brand/framatex-app.css',
                     '/assets/brand/fonts/dm-sans-latin-400-normal.woff2', '/assets/landing/redesign/espace-client.webp',
                     '/assets/landing/value-chain/study-needs.png', '/assets/landing/redesign/notre-solution.webp', '/assets/client/client.js'):
            assert anon.get(path).status_code == 200, path
    qa.run('Landing', 'Logo officiel, design system, polices et visuels servis localement', assets)
    def landing_redirect():
        script = anon.get('/assets/landing/landing-redirect.js').text
        assert "'/espace-framatex' + hash" in script and '#\\/launch\\/' in script
    qa.run('QR', 'Anciens liens /#/launch/<id> redirigés vers /espace-framatex', landing_redirect)

    # ------------------------------------------------------------- Sécurité interne
    def internal_requires_session():
        for path in ('/api/workspace/dashboard', '/api/workspace/search', '/api/workspace/activity', '/api/launches', '/api/admin/clients'):
            assert anon.get(path).status_code == 401, path
    qa.run('Sécurité', 'Routes internes : session FRAMATEX obligatoire (401)', internal_requires_session)
    def csrf_header():
        r = anon.post('/api/auth/login', json={'email': 'x@y.z', 'password': 'x'})
        assert r.status_code == 403, r.status_code
    qa.run('Sécurité', 'Écriture sans en-tête X-Framatex-Request refusée (403)', csrf_header)
    login = admin.post('/api/auth/login', json={'email': 'recette@framatex.test', 'password': password})
    qa.check('FRAMATEX', 'Connexion interne email + mot de passe', login.status_code == 200, login.text)
    me = admin.get('/api/auth/me')
    qa.check('FRAMATEX', 'GET /api/auth/me après connexion : profil administrateur',
             me.status_code == 200 and me.json().get('admin') is True and me.json().get('email') == 'recette@framatex.test', me.text[:200])
    bad = anon.post('/api/auth/login', headers=H, json={'email': 'recette@framatex.test', 'password': 'mauvais-mot-de-passe'})
    qa.check('FRAMATEX', 'Connexion interne refusée avec mauvais mot de passe (401)', bad.status_code == 401)

    launches = {l['order_ref']: l for l in admin.get('/api/launches').json()}
    A, B, C, D, E = (launches[f'FR-DEMO-2026-014{i}']['id'] for i in range(1, 6))
    def four_folders():
        for code, launch_id in zip('ABCDE', (A, B, C, D, E)):
            folders = admin.get(f'/api/launches/{launch_id}').json()['folders']
            assert [f['slug'] for f in folders] == ['jaune', 'rose', 'bleu', 'vert'], (code, folders)
    qa.run('Dossier Jaune', 'Chaque dossier A–E possède ses 4 sous-dossiers JAUNE, ROSE, BLEU, VERT', four_folders)

    # ------------------------------------------------------------- Dashboard / complétude / manquants
    dash = admin.get('/api/workspace/dashboard').json()
    def kpis():
        k = dash['kpis']; rows = dash['launches']
        assert k and all(type(value) is int and value >= 0 for value in k.values()), k
        assert k['active_launches'] == len(rows) == 5
        assert k['complete_launches'] == sum(r['completion']['state'] == 'complete' for r in rows)
        assert k['incomplete_launches'] == sum(r['completion']['state'] == 'incomplete' for r in rows)
        assert k['missing_documents'] == len(dash['missing_documents'])
        assert k['pending_documents'] > 0 and k['alerts'] > 0 and k['complete_launches'] > 0
    qa.run('Dashboard', 'KPI calculés à partir des données réelles', kpis)
    def missing_docs():
        missing = {(m['order_ref'], m['folder'], m['title']) for m in dash['missing_documents']}
        assert ('FR-DEMO-2026-0142', 'jaune', 'Calcul du besoin') in missing, missing
        assert ('FR-DEMO-2026-0143', 'bleu', 'Pièce douanière') in missing, missing
        assert all(m['status'] == 'missing' and m['due'] for m in dash['missing_documents'])
        assert all(m['launch_id'] == launches[m['order_ref']]['id'] for m in dash['missing_documents'])
    qa.run('Complétude', 'Documents manquants : calcul du besoin (JAUNE, B) et pièce douanière (BLEU, C), lien vers la bonne commande', missing_docs)
    qa.run('Dashboard', 'Actions prioritaires présentes et rattachées à un dossier', lambda: bool(dash['priorities']['items']) and all(p.get('launch_id') for p in dash['priorities']['items']))
    qa.run('Dashboard', 'Dernières activités issues de l’audit', lambda: bool(dash['activity']) and all(a.get('action') and a.get('actor') for a in dash['activity']))

    statuses = {}
    for ref, launch in launches.items():
        comp = admin.get(f"/api/launches/{launch['id']}/completeness").json()
        for row in comp['rows']:
            statuses.setdefault(row['status'], []).append((ref, row['title'], row['due']))
    def not_due_vs_missing():
        # C (étape expédition) : la pièce douanière manque. Revenir à l'étape « approvisionnement »
        # la rend non exigible : elle doit passer en not_due et quitter la liste des documents manquants.
        comp = admin.get(f'/api/launches/{C}/completeness').json()
        douane = next(r for r in comp['rows'] if r['title'] == 'Pièce douanière')
        assert douane['status'] == 'missing' and comp['phase'] == 'expedition', (douane['status'], comp['phase'])
        r = admin.patch(f'/api/launches/{C}/completeness/context', json={'profile_id': comp['profile_id'], 'phase': 'approvisionnement'})
        assert r.status_code == 200, r.text
        comp2 = admin.get(f'/api/launches/{C}/completeness').json()
        row = next(r for r in comp2['rows'] if r['title'] == 'Pièce douanière')
        assert row['status'] == 'not_due' and not row['due'], row['status']
        dash2 = admin.get('/api/workspace/dashboard').json()
        assert not any(m['launch_id'] == C and m['title'] == 'Pièce douanière' for m in dash2['missing_documents']), 'pièce non exigible comptée manquante'
        admin.patch(f'/api/launches/{C}/completeness/context', json={'profile_id': comp['profile_id'], 'phase': 'expedition'})
        statuses.setdefault('not_due', []).append(('C', row['title'], False))
    qa.run('Complétude', 'État not_due distinct de missing (jamais compté manquant)', not_due_vs_missing)
    def exception_state():
        comp = admin.get(f'/api/launches/{C}/completeness').json()
        row = next(r for r in comp['rows'] if r['status'] == 'missing')
        r = admin.put(f"/api/launches/{C}/completeness/rules/{row['id']}/exception", json={'reason': 'Dérogation douanière accordée pour ce lot (recette).', 'active': True})
        assert r.status_code == 200, r.text
        comp = admin.get(f'/api/launches/{C}/completeness').json()
        assert next(x for x in comp['rows'] if x['id'] == row['id'])['status'] == 'exception'
        statuses.setdefault('exception', []).append(('C', row['title'], True))
        admin.put(f"/api/launches/{C}/completeness/rules/{row['id']}/exception", json={'reason': 'Fin de dérogation (recette).', 'active': False})
    qa.run('Complétude', 'État exception (dérogation motivée)', exception_state)
    qa.check('Complétude', 'Les cinq états rencontrés : present, missing, pending_validation, not_due, exception',
             {'present', 'missing', 'pending_validation', 'not_due', 'exception'} <= set(statuses), sorted(statuses))

    # ------------------------------------------------------------- Documents, formats, SHA-256, doublons
    def upload(launch, folder, title, doc_type, name, data, comment=''):
        return admin.post(f'/api/launches/{launch}/documents', data={'folder': folder, 'title': title, 'doc_type': doc_type, 'comment': comment},
                          files={'file': (name, data, 'application/octet-stream')})
    samples = {'pdf': ('jaune', 'Rapport de contrôle', 'controle.pdf', pdf_bytes()),
               'docx': ('jaune', 'Dossier technique client', 'specification.docx', docx_bytes('Spécification recette')),
               'xlsx': ('jaune', 'Calcul du besoin', 'besoin-formats.xlsx', xlsx_bytes({'A1': 10, 'B1': '=A1*2'})),
               'png': ('rose', 'Bon de réception', 'reception.png', image_bytes('PNG')),
               'jpg': ('rose', 'Bon de réception', 'reception.jpg', image_bytes('JPEG')),
               'txt': ('jaune', 'Fiche d’emploi', 'note.txt', 'Note de recette'.encode()),
               'eml': ('rose', 'Courriel', 'message.eml', eml_bytes()),
               'csv': ('vert', 'Facture', 'export.csv', 'ref;quantite\nA;10\n'.encode())}
    uploaded = {}
    for ext, (folder, doc_type, name, data) in samples.items():
        r = upload(A, folder, f'Recette format {ext.upper()}', doc_type, name, data)
        ok = r.status_code == 201 and r.json()['latest']['sha256'] == sha(data)
        qa.check('Documents', f'Dépôt {ext.upper()} accepté avec SHA-256 exact', ok, r.text[:200])
        if r.status_code == 201:
            uploaded[ext] = r.json()
    qa.check('Documents', 'Exécutable refusé (422)', upload(A, 'jaune', 'Exe', 'Autre', 'outil.exe', b'MZ\x90\x00').status_code == 422)
    qa.check('Documents', 'Faux PDF refusé (contrôle de signature, 422)', upload(A, 'jaune', 'Faux', 'Autre', 'faux.pdf', b'pas un pdf').status_code == 422)
    qa.check('Documents', 'Fichier > 20 Mo refusé (413)', upload(A, 'jaune', 'Gros', 'Autre', 'gros.txt', b'x' * (20 * 1024 * 1024 + 1)).status_code == 413)
    qa.check('Documents', 'Sous-dossier invalide refusé (422)', upload(A, 'violet', 'X', 'Autre', 'x.txt', b'x').status_code == 422)
    dup = upload(A, 'jaune', 'Copie de la note', 'Autre', 'copie.txt', 'Note de recette'.encode())
    qa.check('Documents', 'Doublon détecté (même SHA-256 dans le dossier)', dup.status_code == 201 and len(dup.json().get('duplicates', [])) >= 1, dup.text[:200])
    def download_exact():
        doc = uploaded['docx']
        r = admin.get(f"/api/documents/{doc['id']}/versions/1/download")
        assert r.status_code == 200 and r.content == samples['docx'][3]
    qa.run('Documents', 'Téléchargement : octets identiques au dépôt', download_exact)

    # ------------------------------------------------------------- Versions V1 → V2 → V3
    def versions():
        v1, v2, v3 = b'Version 1 qualite', b'Version 2 qualite', b'Version 3 qualite'
        doc = upload(A, 'jaune', 'Fiche qualité recette', 'Fiche qualité', 'qualite-v1.txt', v1).json()
        for i, blob in ((2, v2), (3, v3)):
            r = admin.post(f"/api/documents/{doc['id']}/versions", files={'file': (f'qualite-v{i}.txt', blob)}, data={'comment': f'Révision {i} de recette'})
            assert r.status_code == 201, r.text
        same = admin.post(f"/api/documents/{doc['id']}/versions", files={'file': ('qualite-copie.txt', v2)})
        assert same.status_code == 409, ('version identique acceptée', same.status_code)
        detail = admin.get(f'/api/launches/{A}').json()
        doc_now = next(d for f in detail['folders'] for d in f.get('documents', []) if d['id'] == doc['id'])
        nums = [v['number'] for v in doc_now['versions']]
        assert nums == [1, 2, 3], nums
        assert all(v['created_at'] and v['uploaded_by_name'] == 'Responsable recette' for v in doc_now['versions'])
        assert admin.get(f"/api/documents/{doc['id']}/versions/1/download").content == v1
        activity = admin.get('/api/workspace/activity', params={'document_id': doc['id']}).json()
        actions = [x['action'] for x in activity['items']]
        assert actions.count('document.version.add') == 2 and 'document.upload' in actions, actions
        comments = admin.get(f"/api/document-versions/{doc_now['versions'][2]['id']}/comments").json()
        assert any('Révision 3' in c.get('body', '') for c in comments), comments
        return True
    qa.run('Versions', 'V1 → V2 → V3 : anciennes versions conservées, auteur, dates, téléchargement V1, refus doublon, audit, commentaires', versions)

    # ------------------------------------------------------------- Alertes d'impact
    def impacts():
        detail = admin.get(f'/api/launches/{B}').json()
        doc = next(d for f in detail['folders'] for d in f.get('documents', []) if d['doc_type'] == 'Fiche d’emploi')
        before = {a['id'] for a in admin.get(f'/api/launches/{B}/impacts').json()}
        r = admin.post(f"/api/documents/{doc['id']}/versions", files={'file': ('fiche-emploi-v2.txt', b'Fiche emploi revisee recette')})
        assert r.status_code == 201, r.text
        alerts = [a for a in admin.get(f'/api/launches/{B}/impacts').json() if a['id'] not in before]
        assert alerts and all(a['status'] == 'to_review' for a in alerts), alerts
        assert any('besoin' in a['instruction'].lower() and 'prix' in a['instruction'].lower() for a in alerts)
        r = admin.post(f"/api/impacts/{alerts[0]['id']}/decision", json={'status': 'acknowledged', 'note': 'Calcul du besoin et fiche de prix revus (recette).'})
        assert r.status_code in (200, 201), r.text
        # Les deux autres décisions sur les alertes existantes de E et C si disponibles
        others = [a for l in (E, D, C, A) for a in admin.get(f'/api/launches/{l}/impacts').json() if a['status'] == 'to_review']
        for alert, status in zip(others, ('action_required', 'not_applicable')):
            rr = admin.post(f"/api/impacts/{alert['id']}/decision", json={'status': status, 'note': f'Décision de recette : {status}.'})
            assert rr.status_code in (200, 201), rr.text
        return True
    qa.run('Alertes', 'Nouvelle version de fiche d’emploi → alerte (besoin + prix), décisions humaines tracées', impacts)

    # ------------------------------------------------------------- Cohérence
    def coherence():
        q = admin.get(f'/api/launches/{E}/quality').json()
        statuses_q = {row['status'] for row in q['rows']}
        assert statuses_q, q
        row = q['rows'][0]
        assert all('decision' not in row or row.get('decision') in (None, {}) or row['status'] for row in q['rows'])
        return f"états observés : {sorted(statuses_q)}"
    qa.run('Cohérence', 'Contrôle de cohérence calculé, sans décision automatique', coherence)
    def coherence_states():
        observed = set()
        for l in (A, B, C, D, E):
            observed |= {row['status'] for row in admin.get(f'/api/launches/{l}/quality').json()['rows']}
        # Unité incompatible : une valeur relevée n'est jamais réécrite, on relève sur une nouvelle version.
        q = admin.get(f'/api/launches/{E}/quality').json()
        row = next(r for r in q['rows'] if r['pair'])
        nv = admin.post(f"/api/documents/{row['pair'][0]['document_id']}/versions", files={'file': ('lancement-provisoire-v2.txt', b'Lancement provisoire revise : 150 cartons')})
        assert nv.status_code == 201, nv.text
        version_id = nv.json()['versions'][-1]['id']
        assert admin.post(f'/api/document-versions/{version_id}/data', json={'field_key': row['field_key'], 'value': '150', 'unit': 'carton'}).status_code == 201
        observed |= {x['status'] for x in admin.get(f'/api/launches/{E}/quality').json()['rows']}
        # Cohérent puis ambigu sur le dossier A : deux pièces égales, puis un second candidat.
        def doc_with_value(doc_type, name, value):
            d = upload(A, 'jaune', f'{doc_type} recette', doc_type, name, f'{doc_type} : {value} pieces'.encode()).json()
            assert admin.post(f"/api/document-versions/{d['versions'][-1]['id']}/data", json={'field_key': row['field_key'], 'value': value, 'unit': 'pièce'}).status_code == 201
        doc_with_value('Lancement provisoire', 'provisoire-a.txt', '1800')
        doc_with_value('Lancement définitif', 'definitif-a.txt', '1800')
        observed |= {x['status'] for x in admin.get(f'/api/launches/{A}/quality').json()['rows']}
        doc_with_value('Lancement définitif', 'definitif-a-bis.txt', '1750')
        observed |= {x['status'] for x in admin.get(f'/api/launches/{A}/quality').json()['rows']}
        expected = {'consistent', 'difference', 'unavailable', 'unit_mismatch', 'ambiguous'}
        assert expected <= observed, f'états observés : {sorted(observed)}'
        return True
    qa.run('Cohérence', 'États cohérent / écart / donnée manquante / unité incompatible / ambigu', coherence_states)
    def quality_decision_requires_human_note():
        q = admin.get(f'/api/launches/{E}/quality').json()
        row = next(r for r in q['rows'] if r['status'] in ('difference', 'unit_mismatch'))
        r = admin.post(f"/api/launches/{E}/quality/{row['rule_id']}/decisions", json={'signature': row.get('signature', '0' * 64), 'decision': 'justified', 'note': 'court'})
        assert r.status_code == 422, r.status_code
    qa.run('Cohérence', 'Décision de cohérence : justification humaine obligatoire', quality_decision_requires_human_note)

    # ------------------------------------------------------------- OCR / classification
    def ocr():
        before = admin.get(f'/api/launches/{B}').json()
        count_before = sum(len(f.get('documents', [])) for f in before['folders'])
        img = image_bytes('PNG', 'CALCUL DU BESOIN\nQuantite : 1800 pieces')
        r = admin.post(f'/api/launches/{B}/intake/preview', files={'file': ('calcul-du-besoin.png', img, 'image/png')})
        assert r.status_code == 200, r.text
        res = r.json()
        assert res['human_validation_required'] is True and res['classification']['needs_confirmation'] is True
        assert res['method'] == 'image_ocr' and '1800' in res['text_excerpt'], res
        assert res['classification']['suggestions'], res
        after = admin.get(f'/api/launches/{B}').json()
        assert sum(len(f.get('documents', [])) for f in after['folders']) == count_before, 'document créé sans confirmation'
        return f"méthode {res['method']}, champs {[f['key'] for f in res['fields']]}, suggestion {res['classification']['suggestions'][0]['doc_type']}"
    qa.run('OCR', 'OCR + classification proposés, confirmation humaine obligatoire, aucun dépôt silencieux', ocr)
    def scanned_pdf():
        from PIL import Image
        page = Image.open(io.BytesIO(image_bytes('PNG', 'FACTURE FOURNISSEUR\nMontant total 1250 MAD'))).convert('RGB')
        out = io.BytesIO(); page.save(out, format='PDF')  # PDF sans couche texte, comme un scan
        r = admin.post(f'/api/launches/{B}/intake/preview', files={'file': ('scan.pdf', out.getvalue(), 'application/pdf')})
        assert r.status_code == 200, r.text
        res = r.json()
        assert res['method'] == 'pdf_ocr', res['method']
        assert 'FACTURE' in res['text_excerpt'].upper(), res['text_excerpt']
        return f"méthode {res['method']}, extrait {res['text_excerpt'][:40]!r}"
    qa.run('OCR', 'PDF scanné (sans texte) : OCR page par page', scanned_pdf)

    # ------------------------------------------------------------- Excel
    def excel():
        base_cells = {'A1': 10, 'A2': 20, 'B1': '=A1*2', 'B2': '=A2*2', 'B3': '=SUM(B1:B2)'}
        doc = upload(D, 'jaune', 'Calcul du besoin Excel recette', 'Calcul du besoin', 'besoin.xlsx', xlsx_bytes(base_cells)).json()
        variants = [('modifiée', {**base_cells, 'B1': '=A1*3'}), ('constante', {**base_cells, 'B1': '=A1*3', 'B2': 40}),
                    ('#REF!', {**base_cells, 'B1': '=A1*3', 'B2': 40, 'B3': '=SUM(#REF!)'})]
        for _, cells in variants:
            r = admin.post(f"/api/documents/{doc['id']}/versions", files={'file': ('besoin.xlsx', xlsx_bytes(cells))})
            assert r.status_code == 201, r.text
        rows = [x for x in admin.get(f'/api/launches/{D}/excel-inspections').json() if x['document_id'] == doc['id']]
        by_version = {x['version']: x['inspection'] for x in rows}
        assert by_version[1]['status'] == 'first_version'
        kinds = {v: {f['type'] for f in by_version[v]['findings']} for v in (2, 3, 4)}
        assert 'formula_changed' in kinds[2], kinds
        assert 'formula_removed' in kinds[3], kinds
        assert 'broken_reference' in kinds[4], kinds
        messages = {f['message'] for v in (2, 3, 4) for f in by_version[v]['findings']}
        return f'messages : {sorted(messages)}'
    qa.run('Excel', 'Formule normale, modifiée, remplacée par une constante, #REF! détectées avec message clair', excel)

    # ------------------------------------------------------------- EML
    def eml():
        r = admin.post(f'/api/launches/{C}/emails/import', files={'file': ('confirmation.eml', eml_bytes(), 'message/rfc822')})
        assert r.status_code == 201, r.text
        emails = admin.get(f'/api/launches/{C}/emails').json()['emails']
        item = next(x for x in emails if 'Confirmation delai' in x['subject'])
        assert 'atelier@exemple.test' in item['sender'] and item.get('sent_at_raw'), item
        sugg = admin.get(f"/api/launches/{C}/emails/{item['id']}/suggestions")
        assert sugg.status_code == 200
        bad = admin.post(f'/api/launches/{C}/emails/import', files={'file': ('pas-un-mail.txt', b'x')})
        assert bad.status_code == 422
        return f"sujet, expéditeur, date et {len(sugg.json()) if isinstance(sugg.json(), list) else 'des'} suggestion(s) d’engagement"
    qa.run('EML', 'Import EML : sujet, expéditeur, date, suggestions d’engagement, refus hors .eml', eml)

    # ------------------------------------------------------------- Recherche
    def search(params, predicate):
        res = admin.get('/api/workspace/search', params=params).json()
        return res['total'] > 0 and all(predicate(x) for x in res['items']), res['total']
    checks = [('commande', {'q': 'FR-DEMO-2026-0143'}, lambda x: x['order_ref'] == 'FR-DEMO-2026-0143'),
              ('client', {'q': 'Collection Atlas'}, lambda x: 'Collection Atlas' in (x['client_name'] or '')),
              ('modèle', {'q': 'Pantalon lin'}, lambda x: x['launch_id'] == C),
              ('référence', {'q': 'LP-DEMO-C'}, lambda x: x['launch_id'] == C),
              ('nom de document', {'q': 'Pièce douanière'}, lambda x: 'pièce douanière' in x['title'].lower()),
              ('type documentaire', {'doc_type': 'Facture'}, lambda x: x['doc_type'] == 'Facture'),
              ('sous-dossier', {'folder': 'bleu'}, lambda x: x['folder'] == 'bleu')]
    for label, params, pred in checks:
        ok, total = search(params, pred)
        qa.check('Recherche', f'Recherche par {label} : résultats pertinents rattachés à leur dossier', ok, f'{total} résultat(s)')

    # ------------------------------------------------------------- Mémoire interne
    secret_note = f'Note interne confidentielle {secrets.token_hex(4)}'
    for party in ('client', 'fournisseur', 'atelier'):
        r = admin.post(f'/api/launches/{A}/experience-notes', json={'party_type': party, 'party_name': f'Partenaire {party}', 'category': 'retour', 'body': f'{secret_note} · {party}'})
        qa.check('Mémoire interne', f'Ajout d’une note {party}', r.status_code == 201, r.text[:150])
    notes = admin.get('/api/workspace/experience').json()
    qa.check('Mémoire interne', 'Lecture des notes Client / Fournisseur / Atelier associées au dossier',
             {n['party_type'] for n in notes if secret_note in n['body'] and n.get('launch_id') == A} == {'client', 'fournisseur', 'atelier'})

    # ------------------------------------------------------------- Client ↔ FRAMATEX
    client_pw = secrets.token_urlsafe(16) + 'Aa1!'   # comptes Client de recette, base temporaire
    qa.secrets.add(client_pw)
    def register(email, company):
        r = anon.post('/api/client/register', headers=H, json={'raison_sociale': company, 'nom_contact': 'Contact recette', 'email': email,
                      'telephone': '+33 6 00 00 00 00', 'adresse': '1 rue de la Recette, 59000 Lille', 'pays': 'France', 'password': client_pw})
        assert r.status_code == 201, r.text
        return r.json()
    acc_a = register('client.a@recette.test', 'Client A Recette')
    acc_b = register('client.b@recette.test', 'Client B Recette')
    qa.check('Client', 'Inscription : compte créé au statut pending', acc_a['status'] == 'pending')
    ca = httpx.Client(base_url=base, headers=H, timeout=30); cb = httpx.Client(base_url=base, headers=H, timeout=30)
    r = ca.post('/api/client/login', json={'email': 'client.a@recette.test', 'password': client_pw})
    qa.check('Client', 'Connexion refusée tant que le compte est pending (403 account_pending)', r.status_code == 403 and 'pending' in r.text, r.text)
    for acc in (acc_a, acc_b):
        assert admin.patch(f"/api/admin/clients/{acc['id']}/status", json={'status': 'active'}).status_code == 200
    qa.check('Client', 'Activation par un administrateur FRAMATEX', True)
    qa.check('Client', 'Connexion Client A après activation', ca.post('/api/client/login', json={'email': 'client.a@recette.test', 'password': client_pw}).status_code == 200)
    cb.post('/api/client/login', json={'email': 'client.b@recette.test', 'password': client_pw})
    qa.check('Client', 'Mauvais mot de passe client : refus générique (401)', anon.post('/api/client/login', headers=H, json={'email': 'client.a@recette.test', 'password': 'faux-mot-de-passe'}).status_code == 401)
    def client_cannot_call_internal():
        for path in ('/api/workspace/dashboard', '/api/workspace/search', '/api/workspace/experience', '/api/launches', '/api/admin/clients', '/api/admin/client-order-requests'):
            assert ca.get(path).status_code == 401, path
    qa.run('Sécurité', 'Session Client : aucune API interne accessible (401)', client_cannot_call_internal)

    # Workflow de demande de commande
    draft = ca.post('/api/client/order-requests', json={'article': 'Chemise recette', 'quantite': 500, 'date_souhaitee': '2027-05-01', 'status': 'draft'}).json()
    sub = ca.patch(f"/api/client/order-requests/{draft['id']}/status", json={'status': 'submitted'})
    qa.check('Workflow', 'Client : brouillon puis envoi (draft → submitted)', draft['status'] == 'draft' and sub.status_code == 200 and sub.json()['status'] == 'submitted', sub.text[:150])
    forbidden = ca.patch(f"/api/client/order-requests/{draft['id']}/status", json={'status': 'accepted'})
    qa.check('Workflow', 'Client ne peut pas décider d’un statut métier (accepted refusé)', forbidden.status_code in (403, 409, 422), forbidden.status_code)
    s1 = admin.patch(f"/api/admin/client-order-requests/{draft['id']}/status", json={'status': 'under_review'})
    s2 = admin.patch(f"/api/admin/client-order-requests/{draft['id']}/status", json={'status': 'accepted'})
    s3 = admin.patch(f"/api/admin/client-order-requests/{draft['id']}/status", json={'status': 'rejected'})
    qa.check('Workflow', 'FRAMATEX : submitted → under_review → accepted ; accepted terminal (409)', s1.status_code == 200 and s2.json()['status'] == 'accepted' and s3.status_code == 409, s3.status_code)
    other = ca.post('/api/client/order-requests', json={'article': 'Robe recette', 'quantite': 120, 'date_souhaitee': '2027-06-01', 'status': 'submitted'}).json()
    rj = admin.patch(f"/api/admin/client-order-requests/{other['id']}/status", json={'status': 'rejected'})
    qa.check('Workflow', 'FRAMATEX : submitted → rejected', rj.status_code == 200 and rj.json()['status'] == 'rejected')
    qa.check('Workflow', 'Statut visible côté Client', ca.get(f"/api/client/order-requests/{draft['id']}").json()['status'] == 'accepted')

    # Pièces demandées : requested → uploaded → under_review → correction_required → uploaded → accepted
    def document_request():
        dr = admin.post(f"/api/admin/client-order-requests/{draft['id']}/document-requests", json={'title': 'Tableau des mesures', 'description': 'Tailles 36 à 46.'}).json()
        assert ca.get('/api/client/document-requests').json()[0]['status'] == 'requested'
        up1 = ca.post(f"/api/client/document-requests/{dr['id']}/upload", files={'file': ('mesures-v1.txt', b'Mesures v1')})
        assert up1.status_code == 201, up1.text
        item = admin.get(f"/api/admin/client-document-requests/{dr['id']}").json()
        assert item['status'] == 'uploaded', item['status']
        v1 = item['document']['versions'][-1]['id']
        assert admin.patch(f"/api/admin/client-document-requests/{dr['id']}/status", json={'status': 'under_review', 'version_id': v1}).json()['status'] == 'under_review'
        no_msg = admin.patch(f"/api/admin/client-document-requests/{dr['id']}/status", json={'status': 'correction_required', 'version_id': v1, 'message': ''})
        assert no_msg.status_code == 422, 'correction sans message acceptée'
        assert admin.patch(f"/api/admin/client-document-requests/{dr['id']}/status", json={'status': 'correction_required', 'version_id': v1, 'message': 'Ajouter les tailles 40 à 46.'}).json()['status'] == 'correction_required'
        client_view = ca.get(f"/api/client/document-requests/{dr['id']}").json()
        assert client_view['status'] == 'correction_required' and any(ev.get('message') == 'Ajouter les tailles 40 à 46.' for ev in client_view['history'])
        assert ca.post(f"/api/client/document-requests/{dr['id']}/upload", files={'file': ('mesures-v2.txt', b'Mesures v2')}).status_code == 201
        item = admin.get(f"/api/admin/client-document-requests/{dr['id']}").json()
        v2 = item['document']['versions'][-1]['id']
        stale = admin.patch(f"/api/admin/client-document-requests/{dr['id']}/status", json={'status': 'accepted', 'version_id': v1})
        assert stale.status_code == 409, 'décision sur une ancienne version acceptée'
        assert admin.patch(f"/api/admin/client-document-requests/{dr['id']}/status", json={'status': 'accepted', 'version_id': v2}).json()['status'] == 'accepted'
        locked = ca.post(f"/api/client/document-requests/{dr['id']}/upload", files={'file': ('mesures-v3.txt', b'Mesures v3')})
        assert locked.status_code == 409, 'dépôt possible après acceptation'
        assert ca.get(f"/api/client/document-requests/{dr['id']}/versions/1/download").content == b'Mesures v1'
        return dr['id']
    dr_id = None
    try:
        dr_id = document_request(); qa.check('Client ↔ FRAMATEX', 'Pièce demandée : À déposer → Déposé → En vérification → À corriger → V2 → Accepté (versions conservées)', True)
    except Exception as exc:
        qa.check('Client ↔ FRAMATEX', 'Pièce demandée : À déposer → Déposé → En vérification → À corriger → V2 → Accepté (versions conservées)', False, exc)

    # Publication FRAMATEX → Client
    share_id = None
    def unpublished_invisible():
        detail = admin.get(f'/api/launches/{A}').json()
        doc = next(d for f in detail['folders'] for d in f.get('documents', []))
        assert ca.get('/api/client/shared-documents').json() == [], 'document visible sans publication'
        for path in (f"/api/documents/{doc['id']}/versions/1/download", f'/api/launches/{A}', '/api/workspace/search?q=FR-DEMO'):
            assert ca.get(path).status_code == 401, path
    qa.run('Sécurité', 'Documents internes non publiés : invisibles et non téléchargeables par un Client', unpublished_invisible)
    def publication():
        nonlocal share_id
        detail = admin.get(f'/api/launches/{A}').json()
        doc = next(d for f in detail['folders'] if f['slug'] == 'jaune' for d in f.get('documents', []) if d['doc_type'] == 'Dossier technique client' and d['title'] != 'Recette format DOCX')
        before_candidates = {c['version_id'] for c in admin.get('/api/admin/client-share-candidates').json()}
        assert doc['latest']['id'] not in before_candidates, 'version non approuvée partageable'
        assert admin.post(f"/api/documents/{doc['id']}/reviews", json={'decision': 'approved', 'comment': 'Validé pour transmission (recette).'}).status_code == 201
        assert doc['latest']['id'] in {c['version_id'] for c in admin.get('/api/admin/client-share-candidates').json()}
        pub = admin.post(f"/api/admin/clients/{acc_a['id']}/shared-documents", json={'document_version_id': doc['latest']['id'], 'order_request_id': draft['id'], 'reason': 'Transmission du dossier technique validé'})
        assert pub.status_code == 201, pub.text
        share_id = pub.json()['id']
        shared = ca.get('/api/client/shared-documents').json()
        assert [s['id'] for s in shared] == [share_id], 'document non partagé visible ou partage absent'
        assert ca.get(f'/api/client/shared-documents/{share_id}/download').status_code == 200
        assert cb.get('/api/client/shared-documents').json() == []
        assert cb.get(f'/api/client/shared-documents/{share_id}/download').status_code == 404
        price = admin.get('/api/admin/client-share-candidates').json()
        assert not any('prix' in c['title'].lower() or 'facture' in c['title'].lower() for c in price), 'document sensible partageable'
        rv = admin.post(f'/api/admin/client-shared-documents/{share_id}/revoke', json={'reason': 'Retrait de recette du partage client'})
        assert rv.status_code == 200, rv.text
        assert ca.get('/api/client/shared-documents').json() == []
        assert ca.get(f'/api/client/shared-documents/{share_id}/download').status_code == 404
        assert admin.get(f"/api/documents/{doc['id']}/versions/{doc['latest']['number']}/download").status_code == 200
    qa.run('Client ↔ FRAMATEX', 'Partage explicite → visible et téléchargeable par A seulement → retrait effectif, document interne conservé, pièces sensibles exclues', publication)

    # Isolation A / B
    def isolation():
        a_req = ca.post('/api/client/order-requests', json={'article': 'Pièce privée A', 'quantite': 10, 'date_souhaitee': '2027-07-01', 'status': 'submitted'}).json()
        doc = ca.post(f"/api/client/order-requests/{a_req['id']}/documents", data={'title': 'Privé A'}, files={'file': ('prive.txt', b'Prive A')}).json()
        checks = [cb.get(f"/api/client/order-requests/{a_req['id']}").status_code,
                  cb.get(f"/api/client/order-requests/{a_req['id']}/documents/{doc['id']}/versions/1/download").status_code,
                  cb.post(f"/api/client/order-requests/{a_req['id']}/documents", data={'title': 'Intrusion'}, files={'file': ('x.txt', b'x')}).status_code,
                  cb.patch(f"/api/client/order-requests/{a_req['id']}/status", json={'status': 'submitted'}).status_code]
        if dr_id:
            checks += [cb.get(f'/api/client/document-requests/{dr_id}').status_code,
                       cb.get(f'/api/client/document-requests/{dr_id}/versions/1/download').status_code]
        assert all(code == 404 for code in checks), checks
        assert all(r['id'] != a_req['id'] for r in cb.get('/api/client/order-requests').json())
        assert cb.get('/api/client/document-requests').json() == []
    qa.run('Sécurité', 'Client B ne voit ni demandes, ni documents, ni pièces demandées, ni partages de A (404)', isolation)
    def order_to_launch():
        order = ca.post('/api/client/order-requests', json={'article': 'Veste recette dossier', 'quantite': 640, 'date_souhaitee': '2027-06-15',
                        'reference_client': 'REF-QA-640', 'status': 'submitted'}).json()
        created = admin.post(f"/api/admin/client-order-requests/{order['id']}/launch")
        assert created.status_code == 201, created.text
        launch_id = created.json()['launch']['id']
        assert created.json()['order']['status'] == 'accepted'
        assert ca.get(f"/api/client/order-requests/{order['id']}").json()['status'] == 'accepted'
        detail = admin.get(f'/api/launches/{launch_id}').json()
        assert [f['slug'] for f in detail['folders']] == ['jaune', 'rose', 'bleu', 'vert'] and detail['order_ref'] == order['request_number']
        ops = admin.get(f'/api/launches/{launch_id}/operations').json()
        assert ops['article'] == 'Veste recette dossier' and ops['quantity'] == 640, ops
        assert ops['client_request']['reference_client'] == 'REF-QA-640' and ops['client_request']['date_souhaitee'] == '2027-06-15'
        again = admin.post(f"/api/admin/client-order-requests/{order['id']}/launch")
        assert again.status_code == 200 and again.json()['launch']['id'] == launch_id, again.text
        assert ca.post(f"/api/admin/client-order-requests/{order['id']}/launch").status_code == 401
        assert admin.get('/api/workspace/activity', params={'action': 'client.order_request.launch_created'}).json()['total'] >= 1
        return f"dossier {detail['order_ref']} créé une seule fois"
    qa.run('Client ↔ FRAMATEX', 'Demande Client → acceptation → Dossier Jaune (4 sous-dossiers, article, quantité, référence, date, lien, audit, sans doublon)', order_to_launch)
    def notes_never_client():
        payloads = [ca.get('/api/client/order-requests').text, ca.get(f"/api/client/order-requests/{draft['id']}").text,
                    ca.get('/api/client/document-requests').text, ca.get('/api/client/shared-documents').text, ca.get('/api/client/me').text]
        assert not any(secret_note in p for p in payloads)
    qa.run('Mémoire interne', 'Notes internes jamais renvoyées par les API Client', notes_never_client)
    def suspended():
        admin.patch(f"/api/admin/clients/{acc_b['id']}/status", json={'status': 'suspended'})
        me = cb.get('/api/client/me'); login_b = anon.post('/api/client/login', headers=H, json={'email': 'client.b@recette.test', 'password': client_pw})
        assert me.status_code == 403 and login_b.status_code == 403 and 'suspended' in login_b.text, (me.status_code, login_b.text)
    qa.run('Sécurité', 'Compte suspendu : session et connexion refusées (403 account_suspended)', suspended)
    qa.check('Client', 'Déconnexion client : session révoquée', ca.post('/api/client/logout').status_code == 200 and ca.get('/api/client/me').status_code == 401)

    # ------------------------------------------------------------- Modules complémentaires
    # Dossier de recette dédié : les dossiers de démonstration A–E ne sont pas modifiés.
    scratch = admin.post('/api/launches', json={'name': 'Recette modules complémentaires', 'order_ref': 'QA-MOD-01',
                                                 'client_name': 'Client recette'}).json()['id']
    def scratch_doc(title, doc_type, text, folder='jaune'):
        r = admin.post(f'/api/launches/{scratch}/documents', data={'folder': folder, 'title': title, 'doc_type': doc_type},
                       files={'file': (f'{title}.txt'.replace(' ', '-'), text.encode(), 'text/plain')})
        assert r.status_code == 201, r.text
        return r.json()
    tech = scratch_doc('Cahier des charges recette', 'Fiche technique', 'Composition tissu : coton 100% obligatoire. Fermeture légèrement plus résistante.')
    proof = scratch_doc('Essai tissu recette', 'Essai laboratoire', 'Tissu livré : coton 100%')
    role = admin.post('/api/admin/roles', json={'name': 'Contrôle recette', 'permissions': [
        {'action': 'launch.read'}, {'action': 'document.read', 'folder': 'jaune'}, {'action': 'document.review', 'folder': 'jaune'}]})
    collab_pw = secrets.token_urlsafe(18) + 'Aa1!'
    qa.secrets.add(collab_pw)
    created = admin.post('/api/admin/users', json={'name': 'Contrôleur recette', 'email': 'controle@framatex.test',
                                                    'password': collab_pw, 'role_ids': [role.json()['id']]})
    qa.check('FRAMATEX', 'Administration : rôle et compte collaborateur créés', role.status_code == 201 and created.status_code == 201,
             role.text + created.text)
    collab = httpx.Client(base_url=base, headers=H, timeout=30)
    qa.check('FRAMATEX', 'Connexion d’un collaborateur non administrateur',
             collab.post('/api/auth/login', json={'email': 'controle@framatex.test', 'password': collab_pw}).status_code == 200)
    def collab_scope():
        for path in ('/api/admin/clients', '/api/admin/users', '/api/admin/client-order-requests', '/api/orli/imports'):
            assert collab.get(path).status_code == 403, path
        r = collab.get(f'/api/launches/{scratch}').json()
        readable = {f['slug'] for f in r['folders'] if f['can_read']}
        assert readable == {'jaune'}, readable
        assert all(not f['documents'] for f in r['folders'] if not f['can_read'])
        hits = collab.get('/api/workspace/search', params={'q': 'recette'}).json()
        assert all(i['folder'] == 'jaune' for i in hits['items']), hits['items']
    qa.run('Sécurité', 'Collaborateur : routes d’administration refusées (403) et seule la rubrique autorisée visible', collab_scope)
    def access_restriction():
        assert admin.put(f'/api/launches/{scratch}/access', json={'restricted': True}).status_code == 200
        assert collab.get(f'/api/launches/{scratch}').status_code == 404
        assert admin.post(f'/api/launches/{scratch}/access/members', json={'user_id': created.json()['id']}).status_code == 201
        assert collab.get(f'/api/launches/{scratch}').status_code == 200
        assert admin.delete(f"/api/launches/{scratch}/access/members/{created.json()['id']}").status_code == 200
        assert collab.get(f'/api/launches/{scratch}').status_code == 404
        assert admin.put(f'/api/launches/{scratch}/access', json={'restricted': False}).status_code == 200
    qa.run('Sécurité', 'Accès restreint par dossier : membre ajouté puis retiré', access_restriction)
    def approval_circuit():
        tpl = admin.post('/api/admin/approval-templates', json={'name': 'Circuit recette', 'folder_slug': 'jaune',
              'document_type': 'Fiche recette', 'stages': [{'name': 'Contrôle', 'reviewer_role_id': role.json()['id']}]})
        assert tpl.status_code == 201, tpl.text
        doc = scratch_doc('Fiche à valider', 'Fiche recette', 'Consommation 1.8 m')
        run = admin.post(f"/api/documents/{doc['id']}/approval-runs")
        assert run.status_code == 201, run.text
        rid = run.json()['id']
        assert admin.post(f'/api/approval-runs/{rid}/decisions', json={'decision': 'approved'}).status_code == 403  # auteur ≠ valideur
        assert collab.get('/api/notifications').json()['unread'] >= 1
        decided = collab.post(f'/api/approval-runs/{rid}/decisions', json={'decision': 'approved', 'comment': 'Conforme'})
        assert decided.status_code == 200, decided.text
        assert admin.get(f'/api/launches/{scratch}/approval-overview').status_code == 200
    qa.run('Workflow', 'Circuit de validation : séparation auteur/valideur, notification, décision', approval_circuit)
    def specifications():
        req = admin.post(f'/api/launches/{scratch}/requirements', json={'source_version_id': tech['latest']['id'],
              'title': 'Composition', 'category': 'Matière', 'description': 'Le tissu doit être en coton 100 %.',
              'source_excerpt': 'Composition tissu : coton 100% obligatoire'})
        assert req.status_code == 201, req.text
        rid = req.json()['id']
        assert admin.post(f'/api/launches/{scratch}/requirements/{rid}/verify', json={'note': 'Vérification sans preuve'}).status_code == 409
        assert admin.post(f'/api/launches/{scratch}/requirements/{rid}/evidence', json={'evidence_version_id': proof['latest']['id']}).status_code == 200
        assert admin.post(f'/api/launches/{scratch}/requirements/{rid}/verify', json={'note': 'Composition vérifiée sur l’essai'}).json()['state'] == 'verified'
        amb = admin.post(f'/api/launches/{scratch}/ambiguities', json={'source_version_id': tech['latest']['id'],
              'excerpt': 'Fermeture légèrement plus résistante', 'kind': 'intensite_imprecise', 'question': 'Quelle résistance mesurable ?'})
        assert amb.status_code == 201, amb.text
        assert admin.post(f"/api/launches/{scratch}/ambiguities/{amb.json()['id']}/resolve",
                          json={'decision': 'resolved', 'note': 'Fermeture B-21 retenue par le responsable'}).json()['status'] == 'resolved'
        assert admin.get(f"/api/document-versions/{tech['latest']['id']}/spec-suggestions").status_code == 200
        assert admin.get(f'/api/launches/{scratch}/specifications').json()['coverage']['verified'] == 1
    qa.run('Workflow', 'Exigences : preuve obligatoire, vérification, ambiguïté résolue avec note', specifications)
    def materials():
        line = admin.post(f'/api/launches/{scratch}/materials', json={'reference': 'TIS-QA', 'label': 'Tissu jersey', 'unit': 'm',
               'quantity': '2700.000', 'source_version_id': tech['latest']['id'], 'note': 'Besoin théorique vérifié'})
        assert line.status_code == 201, line.text
        lid = admin.get(f'/api/launches/{scratch}/materials').json()['lines'][0]['id']
        mv = admin.post(f'/api/launches/{scratch}/materials/{lid}/movements', json={'kind': 'ordered', 'quantity': '1000',
             'source_version_id': tech['latest']['id'], 'note': 'Bon de commande examiné'})
        assert mv.status_code == 201, mv.text
        assert admin.post(f'/api/launches/{scratch}/materials/{lid}/movements', json={'kind': 'stolen', 'quantity': '1',
                          'source_version_id': tech['latest']['id'], 'note': 'Nature inconnue'}).status_code == 422
    qa.run('Workflow', 'Matières : besoin documenté, mouvement commandé, nature inconnue refusée', materials)
    def impact_simulation():
        changes = {'title': 'Recette', 'changes': {'consommation_matiere': '1.9'}}
        refused = admin.post(f'/api/launches/{scratch}/impact-simulations', json=changes)
        assert refused.status_code == 422, refused.text  # valeurs de référence absentes : aucun calcul inventé
        src = scratch_doc('Fiche emploi simulation', 'Fiche d’emploi', 'Valeurs de référence')['latest']['id']
        for key, value, unit in (('quantite_lancement', '1500', 'pièce'), ('consommation_matiere', '1.8', 'm/pièce'),
                                 ('cout_matiere_unitaire', '25', 'MAD/m')):
            r = admin.post(f'/api/launches/{scratch}/values', json={'key': key, 'value': value, 'unit': unit, 'source_version_id': src})
            assert r.status_code == 201, r.text
        assert len(admin.get(f'/api/launches/{scratch}/dependency-graph').json()['nodes']) == 5
        sim = admin.post(f'/api/launches/{scratch}/impact-simulations', json=changes)
        assert sim.status_code == 200 and sim.json()['saved'] is False, sim.text
        outputs = {o['key']: o for o in sim.json()['result']['outputs']}
        assert outputs and admin.get(f'/api/launches/{scratch}/impact-simulations').json() == []
        assert admin.get(f'/api/launches/{scratch}/values').json()  # valeurs enregistrées inchangées
    qa.run('Alertes', 'Simulation d’impact non destructive (rien n’est enregistré sans demande)', impact_simulation)
    def orli_file_import():
        mapping = {'orli_ref': 'Lancement', 'order_ref': 'Commande', 'client_name': 'Client', 'designation': 'Article', 'quantity': 'Quantite'}
        raw = 'Lancement;Commande;Client;Article;Quantite\nOR-QA-1;QA-MOD-01;Client recette;Vestes;1500\n'
        r = admin.post('/api/orli/imports', data={'mapping': json.dumps(mapping)}, files={'file': ('export_orli.csv', raw.encode(), 'text/csv')})
        assert r.status_code == 201 and r.json()['valid_rows'] == 1, r.text
        assert admin.get(f'/api/launches/{scratch}').json()['orli_ref'] is None  # aucun rapprochement implicite
        assert admin.get(f'/api/launches/{scratch}/orli-reconciliation').json()['candidates']
    qa.run('Excel', 'Import d’un export ORLI (CSV) en lecture seule, sans rapprochement automatique', orli_file_import)
    def supplier_requests():
        r = admin.post(f'/api/launches/{scratch}/supplier-requests', json={'folder_slug': 'rose', 'supplier_name': 'Fournisseur fictif',
            'requested_title': 'Bon de livraison', 'expected_doc_type': 'Bon de livraison', 'due_date': '2026-12-01'})
        assert r.status_code == 201, r.text
        assert admin.get(f'/api/launches/{scratch}/supplier-requests').json()['items']
        assert admin.get(f'/api/launches/{A}/client-preview').status_code == 200
    qa.run('Client ↔ FRAMATEX', 'Demande fournisseur suivie en interne (aucun email envoyé) et aperçu client simulé', supplier_requests)
    def pilot_reports():
        rep = admin.post('/api/pilot/reports', json={'title': 'Recette : libellé à vérifier', 'module': 'Documents et versions',
              'steps': '1. Ouvrir le dossier. 2. Déposer une pièce.', 'expected': 'Pièce visible.', 'observed': 'Pièce visible.',
              'severity': 'mineur', 'launch_id': scratch})
        assert rep.status_code == 201, rep.text
        assert admin.patch(f"/api/pilot/reports/{rep.json()['id']}", json={'status': 'resolu', 'resolution_note': 'Contrôlé pendant la recette finale.'}).status_code == 200
    qa.run('FRAMATEX', 'Retours de recette pilote : signalement puis résolution', pilot_reports)
    def guide():
        known = admin.post('/api/assistant/guide', json={'question': 'Où déposer un document ?', 'launch_id': A})
        unknown = admin.post('/api/assistant/guide', json={'question': 'zzz qwerty'})
        assert known.status_code == 200 and known.json()['answer'], known.text
        assert unknown.json()['title'] == 'Question à préciser', unknown.text
    qa.run('FRAMATEX', 'Framatex Assist : guide de procédures déterministe (aucune IA)', guide)
    def other_views():
        for path in ('/api/dashboard', '/api/workspace/priorities', f'/api/launches/{A}/workflow-summary', f'/api/launches/{A}/reviews',
                     f'/api/launches/{A}/values', '/api/business-value-fields', '/api/quality/fields', '/api/admin/users',
                     '/api/admin/roles', '/api/admin/audit', '/api/workspace/experience', '/api/workspace/collaborators',
                     '/api/workspace/document-types', f'/api/launches/{A}/operations', f'/api/launches/{A}/client-publications'):
            r = admin.get(path); assert r.status_code == 200, (path, r.status_code)
    qa.run('FRAMATEX', 'Vues de pilotage, valeurs, revues, administration et publications : réponses 200', other_views)

    # ------------------------------------------------------------- Audit
    def audit_trail():
        items = admin.get('/api/workspace/activity', params={'limit': 200}).json()['items']
        seen = {i['action'] for i in items}
        wanted = {'session.login', 'client.order_request.launch_created', 'document.upload', 'document.version.add', 'document.review', 'document.comment', 'impact.decision',
                  'completeness.exception.set', 'experience.note.create', 'client.register', 'client.status.update',
                  'client.order_request.status_changed', 'client.document_request.correction_required',
                  'client.document_request.accepted', 'client.shared_document.published', 'client.shared_document.revoked'}
        total = admin.get('/api/workspace/activity').json()['total']
        missing = wanted - seen
        if missing:
            for action in sorted(missing):
                if admin.get('/api/workspace/activity', params={'action': action}).json()['total']:
                    missing.discard(action)
        assert not missing, sorted(missing)
        sample = admin.get('/api/workspace/activity', params={'action': 'document.version.add'}).json()['items'][0]
        assert sample['actor'] and sample['occurred_at'] and sample['target_id'], sample
        return f'{total} opérations journalisées'
    qa.run('Audit', 'Dépôt, version, validation, commentaire, décisions, statuts, publication, retrait journalisés (auteur, date, action, objet)', audit_trail)

    # ------------------------------------------------------------- QR
    def qr():
        link = admin.get(f'/api/launches/{B}/paper-link').json()
        assert link['url'].endswith(f'/#/launch/{B}'), link
        assert admin.get(f'/api/launches/{B}/paper-qr.png').headers['content-type'] == 'image/png'
    qa.run('QR', 'Étiquette QR : lien /#/launch/<id> et image PNG générés', qr)


if __name__ == '__main__':
    main()
